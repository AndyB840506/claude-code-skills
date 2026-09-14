# -*- coding: utf-8 -*-
"""Convierte un SRT de transcriptor (WhisperX + diarizacion) en un .docx presentable.

Reusa el idioma de parsing de SRT de
C:\\Users\\andre\\.claude\\skills\\mrputridsden-production\\wpm.py (regex de timestamp,
split por linea en blanco, strip de etiquetas [SPEAKER_XX]:).

Modos:
  --list-speakers   solo detecta speakers distintos y muestra una linea de ejemplo por cada uno
  (default)         genera el .docx

Uso:
  python srt_to_docx.py <srt_path> --list-speakers
  python srt_to_docx.py <srt_path> <output_docx> --title "Nombre del show" ^
      --speakers SPEAKER_00=S1,SPEAKER_01=S2
"""
import re
import io
import sys
import argparse
from datetime import date

TS = re.compile(r"(\d\d):(\d\d):(\d\d),(\d\d\d)\s*-->\s*(\d\d):(\d\d):(\d\d),(\d\d\d)")
SPEAKER_TAG = re.compile(r"\[SPEAKER_[^\]]*\]\s*:?")
SPEAKER_ID = re.compile(r"\[(SPEAKER_[^\]]*)\]")

MARKER_INTERVAL_SECONDS = 120


def secs(h, m, s, ms):
    return int(h) * 3600 + int(m) * 60 + int(s) + int(ms) / 1000.0


def read_srt(path):
    try:
        return io.open(path, encoding="utf-8").read()
    except UnicodeDecodeError:
        return io.open(path, encoding="latin-1").read()


def parse_segments(raw):
    """Devuelve lista de (start_sec, end_sec, speaker_id_or_None, texto_limpio)."""
    segments = []
    blocks = re.split(r"\n\s*\n", raw)
    for b in blocks:
        m = TS.search(b)
        if not m:
            continue
        start = secs(*m.groups()[:4])
        end = secs(*m.groups()[4:])
        text = TS.sub("", b)
        text = re.sub(r"^\s*\d+\s*$", "", text, flags=re.M)
        text = re.sub(r"<[^>]+>", " ", text)
        spk_match = SPEAKER_ID.search(text)
        speaker_id = spk_match.group(1) if spk_match else None
        text = SPEAKER_TAG.sub("", text)
        text = " ".join(text.split())
        if not text:
            continue
        segments.append((start, end, speaker_id, text))
    return segments


def detect_speakers(segments):
    """Devuelve lista de (speaker_id, primer_texto_de_ejemplo) en orden de aparicion."""
    seen = []
    for _, _, spk, text in segments:
        if spk is None:
            continue
        if spk not in [s for s, _ in seen]:
            seen.append((spk, text))
    return seen


def build_turns(segments, speaker_map):
    """Fusiona segmentos consecutivos del mismo speaker en turnos."""
    turns = []
    for start, end, spk, text in segments:
        label = speaker_map.get(spk, spk or "")
        if turns and turns[-1][2] == label:
            turns[-1][3] += " " + text
            turns[-1][1] = end
        else:
            turns.append([start, end, label, text])
    return turns


def format_ts(seconds):
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return "%02d:%02d:%02d" % (h, m, s)


def parse_speaker_map(arg):
    mapping = {}
    if not arg:
        return mapping
    for pair in arg.split(","):
        if "=" not in pair:
            continue
        k, v = pair.split("=", 1)
        mapping[k.strip()] = v.strip()
    return mapping


def default_speaker_map(detected):
    return {spk: "S%d" % (i + 1) for i, (spk, _) in enumerate(detected)}


def write_docx(turns, output_path, title):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Sample Transcript — %s" % date.today().isoformat())

    last_marker = None
    for start, _end, label, text in turns:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)

        show_marker = last_marker is None or (start - last_marker) >= MARKER_INTERVAL_SECONDS
        prefix = ("[%s] " % format_ts(start)) if show_marker else ""
        if show_marker:
            last_marker = start

        run = p.add_run("%s%s: " % (prefix, label))
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Calibri"

        run2 = p.add_run(text)
        run2.font.size = Pt(11)
        run2.font.name = "Calibri"

    doc.save(output_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("srt_path")
    ap.add_argument("output_docx", nargs="?")
    ap.add_argument("--title", default="Transcript Sample")
    ap.add_argument("--speakers", default="")
    ap.add_argument("--list-speakers", action="store_true")
    args = ap.parse_args()

    raw = read_srt(args.srt_path)
    segments = parse_segments(raw)

    if args.list_speakers:
        detected = detect_speakers(segments)
        print("%d speaker(s) detectado(s):" % len(detected))
        for spk, sample in detected:
            print("  %s: %s" % (spk, sample[:80]))
        return

    if not args.output_docx:
        print("ERROR: falta output_docx (o usa --list-speakers)", file=sys.stderr)
        sys.exit(1)

    detected = detect_speakers(segments)
    speaker_map = default_speaker_map(detected)
    speaker_map.update(parse_speaker_map(args.speakers))

    turns = build_turns(segments, speaker_map)
    write_docx(turns, args.output_docx, args.title)
    print("OK: %s (%d turnos, %d speakers)" % (args.output_docx, len(turns), len(detected)))


if __name__ == "__main__":
    main()
